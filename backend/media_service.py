"""Voice transcription (Whisper), image vision/OCR, and document text extraction.
All processing server-side via the Emergent universal key. Never exposed to the app."""
import os
import io
import base64
import logging
import tempfile
from pathlib import Path

logger = logging.getLogger(__name__)
EMERGENT_LLM_KEY = os.environ["EMERGENT_LLM_KEY"]


async def transcribe_audio(audio_bytes: bytes, filename: str, language: str = "en") -> str:
    from emergentintegrations.llm.openai import OpenAISpeechToText
    suffix = Path(filename).suffix.lower() or ".m4a"
    if suffix not in {".m4a", ".mp3", ".wav", ".webm", ".mp4", ".mpeg", ".mpga", ".aac", ".ogg"}:
        suffix = ".m4a"
    tmp_name = None
    try:
        with tempfile.NamedTemporaryFile(suffix=suffix, delete=False) as tmp:
            tmp.write(audio_bytes)
            tmp_name = tmp.name
        stt = OpenAISpeechToText(api_key=EMERGENT_LLM_KEY)
        kwargs: dict = {"model": "whisper-1"}
        # "auto" (or unknown) -> let Whisper detect the language (English/Hindi/Hinglish etc.)
        if language in ("en", "hi"):
            kwargs["language"] = language
        # The underlying OpenAI client needs a real file object (a bare path string is rejected).
        with open(tmp_name, "rb") as fh:
            result = await stt.transcribe(fh, **kwargs)
        text = result if isinstance(result, str) else getattr(result, "text", None)
        if text is None and isinstance(result, dict):
            text = result.get("text", "")
        return (text or "").strip()
    finally:
        if tmp_name:
            Path(tmp_name).unlink(missing_ok=True)


async def image_qa(image_bytes: bytes, mime_type: str, question: str) -> str:
    from emergentintegrations.llm.chat import LlmChat, UserMessage, ImageContent
    from uuid import uuid4
    image_b64 = base64.b64encode(image_bytes).decode("ascii")
    chat = LlmChat(
        api_key=EMERGENT_LLM_KEY,
        session_id=f"image-{uuid4()}",
        system_message=(
            "You analyze user-provided images (screenshots, invoices, photos). Read visible text "
            "carefully with OCR. If text is unclear, say so; never invent missing values. "
            "You understand English, Hindi and Hinglish."
        ),
    ).with_model("openai", "gpt-4o")
    msg = UserMessage(text=question, file_contents=[ImageContent(image_b64)])
    result = await chat.send_message(msg)
    return result if isinstance(result, str) else getattr(result, "text", str(result))


def extract_document_text(data: bytes, filename: str, mime: str) -> str:
    """Best-effort text extraction from common document formats."""
    ext = Path(filename).suffix.lower()
    try:
        if ext == ".pdf" or mime == "application/pdf":
            from pypdf import PdfReader
            reader = PdfReader(io.BytesIO(data))
            return "\n".join((p.extract_text() or "") for p in reader.pages[:40]).strip()
        if ext in (".docx",):
            import docx
            d = docx.Document(io.BytesIO(data))
            return "\n".join(p.text for p in d.paragraphs).strip()
        if ext in (".xlsx",):
            import openpyxl
            wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
            out = []
            for ws in wb.worksheets[:5]:
                out.append(f"# Sheet: {ws.title}")
                for row in ws.iter_rows(values_only=True):
                    cells = [str(c) for c in row if c is not None]
                    if cells:
                        out.append(" | ".join(cells))
            return "\n".join(out[:1000]).strip()
        if ext in (".csv", ".txt", ".md", ".json"):
            return data.decode("utf-8", errors="ignore")[:60000].strip()
    except Exception as e:
        logger.warning(f"Text extraction failed for {filename}: {e}")
    return ""
